"""
=============================================================================
KNOWLEDGE GRAPH & INFORMATION RETRIEVAL SYSTEMS (KGIRS) VIRTUAL LAB
=============================================================================
Experiment: Relationship Extraction from Text
Objective:  Extract semantic relationships between identified entities from
            natural language text and represent them as entity–relationship
            triples suitable for knowledge graph construction.
Architecture: 4-Section Modular Virtual Laboratory
  1. Theory: Concepts, NER, Relationship Extraction, Pipeline & Procedure
  2. Simulation: Interactive Extraction Sandbox, Triples, Knowledge Graph
  3. Quiz: Self-grading conceptual assessment with instant feedback
  4. Report Generation: Student info, extracted data, observations & PDF export
=============================================================================
"""

import os
import io
import re
import json
import html
import tempfile
import random
from datetime import datetime
from typing import List, Dict, Tuple, Any, Optional

import streamlit as st
import pandas as pd

# Core NLP & Graph Libraries
try:
    import spacy
    from spacy.tokens import Doc, Span, Token
    SPACY_AVAILABLE = True
except ImportError:
    SPACY_AVAILABLE = False

try:
    import networkx as nx
    NETWORKX_AVAILABLE = True
except ImportError:
    NETWORKX_AVAILABLE = False

try:
    from pyvis.network import Network
    import streamlit.components.v1 as components
    PYVIS_AVAILABLE = True
except ImportError:
    PYVIS_AVAILABLE = False

try:
    import matplotlib.pyplot as plt
    MATPLOTLIB_AVAILABLE = True
except ImportError:
    MATPLOTLIB_AVAILABLE = False

# Document Parsers
try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    try:
        import PyPDF2 as pypdf
        PYPDF_AVAILABLE = True
    except ImportError:
        PYPDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# PDF Generator
try:
    from fpdf import FPDF
    FPDF_AVAILABLE = True
except ImportError:
    FPDF_AVAILABLE = False


# =============================================================================
# 1. EXPERIMENT CONFIGURATION & EDUCATIONAL CONTENT
# =============================================================================

EXPERIMENT_CONFIG = {
    "title": "Relationship Extraction from Text",
    "course": "Knowledge Graph & Information Retrieval Systems",
    "objectives": [
        "Understand the principles of Named Entity Recognition (NER) and syntactic dependency parsing for relation discovery.",
        "Extract semantic relationships between entity spans from unstructured text into structured (Subject, Relationship, Object) triples.",
        "Handle syntactic variations including active voice, passive voice inversion, and prepositional attachments.",
        "Construct and visualize an interactive Knowledge Graph from extracted semantic triples.",
        "Validate extracted relationships and export structured triples in CSV and JSON formats."
    ]
}

THEORY_CONTENT = {
    "background": """
### 1. Introduction to Relationship Extraction
**Relationship Extraction (RE)** is a core Information Extraction (IE) task that identifies semantic connections between named entities mentioned in natural language text. In information retrieval and semantic web systems, turning unstructured text into structured facts powers question answering, search engines, and domain-specific knowledge bases.

### 2. The Extraction Pipeline
The transition from raw unstructured text to a structured graph follows a five-stage pipeline:
```
Text  -->  Entities (NER)  -->  Relationships (Dependency Parsing)  -->  Triples (S, R, O)  -->  Knowledge Graph
```

1. **Text Preprocessing**: Tokenization, sentence boundary detection, and normalization of compound terms (e.g., *co-founded* $\\rightarrow$ *co_founded*).
2. **Named Entity Recognition (NER)**: Identifying entity spans and classifying them into semantic categories such as `PERSON`, `ORG` (Organization), `GPE` (Geopolitical Entity), `LOC` (Location), and `PRODUCT`.
3. **Syntactic Dependency Parsing**: Analyzing the grammatical dependency tree to trace connections between subject noun phrases, verbal predicates, and direct/prepositional objects.
4. **Triple Generation**: Structuring extracted relationships into formal **Subject–Relationship–Object** (SPO) triples:
   - $(e_1, r, e_2) \\in \\mathcal{E} \\times \\mathcal{R} \\times \\mathcal{E}$
5. **Knowledge Graph Construction**: Mapping subjects and objects to graph nodes (vertices) and relationships to directed, labeled edges.

### 3. Syntactic & Linguistic Handling
- **Active Voice**: Identifies nominal subjects (`nsubj`) and direct objects (`dobj`/`obj`) linked by a root verb (e.g., *"Steve Jobs founded Apple"* $\\rightarrow$ `Steve Jobs | founded | Apple`).
- **Passive Voice Inversion**: Identifies passive nominal subjects (`nsubjpass`) and agent prepositional phrases (`agent` / `prep` *"by"*) to correctly assign the agent as the semantic subject (e.g., *"Apple was founded by Steve Jobs"* $\\rightarrow$ `Steve Jobs | founded | Apple`).
- **Role & Attribute Structures**: Extracts role affiliations such as *"as CEO of <Org>"* into structured relations (e.g., *"Tim Cook succeeded Steve Jobs as CEO of Apple"* $\\rightarrow$ `Tim Cook | succeeded | Steve Jobs` and `Tim Cook | CEO_of | Apple`).
- **Prepositional Attachments & Comitatives**: Correctly associates location/partnership attachments while distinguishing accompaniment (`with`) from direct targets.
- **Negation Filtering**: Filters out negated statements (e.g., *"Google was acquired by no company"*) to prevent false assertions from entering the knowledge base.
    """,
    "procedure": [
        "Step 1: Review the theoretical background, pipeline stages, and linguistic concepts in this section.",
        "Step 2: Navigate to the Simulation tab in the sidebar menu.",
        "Step 3: Select a sample example, enter custom text, or upload a document (.txt, .pdf, .docx).",
        "Step 4: Click 'Extract Relationships' to trigger NER, dependency parsing, and triple generation.",
        "Step 5: Inspect the detected entity tags, extracted relationship triples table, and interactive Knowledge Graph.",
        "Step 6: Download the extracted triples as CSV or JSON for offline verification.",
        "Step 7: Proceed to the Quiz section to test your understanding of relationship extraction and knowledge graphs.",
        "Step 8: Open Report Generation, enter your student information and observations, and download your official Lab Report PDF."
    ],
    "key_terms": {
        "Named Entity Recognition (NER)": "Subtask of IE that locates and classifies named entities in unstructured text into predefined categories (e.g., PERSON, ORG, GPE).",
        "Relationship Extraction (RE)": "The process of detecting and classifying semantic assertions linking two or more entities.",
        "Semantic Triple (S, R, O)": "A fundamental unit of knowledge representation consisting of Subject (entity), Predicate/Relationship, and Object (entity).",
        "Dependency Parsing": "Analyzing grammatical sentence structure to establish directed head-dependent relationships between words.",
        "Passive Voice Inversion": "Transforming passive syntactic structures ('X was founded by Y') into standard canonical active triples (Y, founded, X).",
        "Knowledge Graph (KG)": "A directed multi-graph where nodes represent entities and edges represent semantic relationships."
    }
}

SAMPLE_TEXTS = {
    "Example 1: Tech Companies & Partnerships": "Steve Wozniak co-founded Apple. Elon Musk founded SpaceX. NASA partnered with SpaceX for space missions.",
    "Example 2: Apple Leadership & Acquisition": "Steve Jobs founded Apple. Steve Wozniak co-founded Apple. Apple acquired Beats Electronics in 2014. Tim Cook succeeded Steve Jobs as CEO of Apple. Apple is headquartered in Cupertino, California.",
    "Example 3: SpaceX & Tesla": "Elon Musk founded SpaceX in 2002. SpaceX is headquartered in Hawthorne, California. Elon Musk also leads Tesla, which manufactures electric vehicles.",
    "Example 4: Scientists & Universities": "Albert Einstein was born in Germany and worked at Princeton University. Marie Curie studied at University of Paris and won the Nobel Prize.",
    "Example 5: Microsoft & GitHub": "Microsoft acquired GitHub in 2018. Satya Nadella is the CEO of Microsoft.",
    "Example 6: Negative Case": "Google was acquired by no company."
}

QUIZ_QUESTIONS = [
    {
        "id": 1,
        "question": "What is the primary objective of Relationship Extraction (RE) in Knowledge Graph construction?",
        "options": [
            "A) To count the total number of words and characters in a paragraph",
            "B) To discover semantic connections between recognized entity spans and represent them as structured triples",
            "C) To compress text documents into binary encodings",
            "D) To translate natural language text into another human language"
        ],
        "answer_index": 1,
        "explanation": "Relationship Extraction identifies and classifies semantic relations between entities, structuring unstructured text into Subject-Relationship-Object triples for knowledge graph ingestion."
    },
    {
        "id": 2,
        "question": "In the semantic triple (Steve Jobs, founded, Apple), which component represents the predicate / relationship?",
        "options": [
            "A) Steve Jobs",
            "B) founded",
            "C) Apple",
            "D) Both Steve Jobs and Apple"
        ],
        "answer_index": 1,
        "explanation": "In standard (Subject, Predicate, Object) knowledge triples, 'founded' is the predicate (directed relation) connecting the subject entity ('Steve Jobs') to the object entity ('Apple')."
    },
    {
        "id": 3,
        "question": "How should an extraction system handle the passive sentence: 'Apple was founded by Steve Jobs'?",
        "options": [
            "A) Discard the sentence because passive verbs cannot express relationships",
            "B) Generate the triple: (Apple, was_founded, Steve Jobs)",
            "C) Invert the agent ('Steve Jobs') to the subject and passive subject ('Apple') to the object: (Steve Jobs, founded, Apple)",
            "D) Duplicate the relationship symmetrically in both directions"
        ],
        "answer_index": 2,
        "explanation": "Passive voice inversion identifies the prepositional agent ('by Steve Jobs') as the true semantic agent (Subject) and canonicalizes the relation to active form ('founded') with 'Apple' as the Object."
    },
    {
        "id": 4,
        "question": "In a Knowledge Graph derived from extracted relationship triples, what do the nodes and directed edges correspond to?",
        "options": [
            "A) Nodes are verbs; Edges are nouns",
            "B) Nodes are sentences; Edges are paragraphs",
            "C) Nodes are entities (concepts/objects); Directed edges are semantic relationships",
            "D) Nodes are characters; Edges are punctuation marks"
        ],
        "answer_index": 2,
        "explanation": "In Knowledge Graph data models (such as RDF and Property Graphs), entities represent the graph nodes (vertices), while semantic relationships form directed, labeled edges."
    },
    {
        "id": 5,
        "question": "Why is Named Entity Recognition (NER) an essential precursor to Relationship Extraction?",
        "options": [
            "A) Because relationship extraction operates by identifying semantic connections specifically between defined entity boundaries",
            "B) Because NER calculates font styling and page layout",
            "C) Because NER removes all verbs from the sentence",
            "D) NER is only used to count total syllables in words"
        ],
        "answer_index": 0,
        "explanation": "NER delineates and classifies the exact entity boundaries (PERSON, ORG, GPE, etc.) that serve as the subject and object arguments for candidate semantic relations."
    },
    {
        "id": 6,
        "question": "In dependency parsing for relationship extraction, what syntactic dependency tag typically identifies the subject of an active-voice clause?",
        "options": [
            "A) dobj (direct object)",
            "B) nsubj (nominal subject)",
            "C) prep (prepositional modifier)",
            "D) punct (punctuation mark)"
        ],
        "answer_index": 1,
        "explanation": "The nsubj (nominal subject) dependency relation identifies the noun phrase that acts as the grammatical subject performing the action in an active sentence."
    },
    {
        "id": 7,
        "question": "What is the purpose of normalizing verbal predicates (e.g., mapping 'co-founded', 'cofounded', and 'co_founded' to a canonical form)?",
        "options": [
            "A) To ensure uniform schema definitions and consistent relation queries in the knowledge graph",
            "B) To delete all verbal assertions from the database",
            "C) To convert numerical dates into Roman numerals",
            "D) To increase the character count of the stored graph"
        ],
        "answer_index": 0,
        "explanation": "Relation canonicalization/normalization aligns syntactic variants onto uniform schema relations, preventing duplicate or fragmented edges in the knowledge graph."
    },
    {
        "id": 8,
        "question": "In the Resource Description Framework (RDF) and property graphs, what is an entity–relationship–entity assertion formally called?",
        "options": [
            "A) A binary byte",
            "B) A semantic triple (Subject, Predicate, Object)",
            "C) A parse token",
            "D) A stopword list"
        ],
        "answer_index": 1,
        "explanation": "Knowledge representation frameworks structure knowledge into atomic Subject-Predicate-Object (SPO) semantic triples."
    },
    {
        "id": 9,
        "question": "How does relationship extraction avoid inserting false facts from negated statements like 'Google was acquired by no company'?",
        "options": [
            "A) By capitalizing all words in the sentence",
            "B) By inspecting syntactic negation modifiers (e.g., neg dependency or negative determiners 'no', 'never') and filtering them out",
            "C) By reversing the direction of the edge in the graph",
            "D) By converting the sentence into passive voice"
        ],
        "answer_index": 1,
        "explanation": "Negation detection checks syntactic negation dependencies and negative quantifiers to prevent false assertions from being populated into the graph."
    },
    {
        "id": 10,
        "question": "What role does coreference resolution play in advanced Relationship Extraction pipelines?",
        "options": [
            "A) It resolves pronouns and aliases (e.g., 'he', 'she', 'it') back to their antecedent entity mentions across sentences",
            "B) It translates text into foreign languages",
            "C) It formats graphs into PDF files",
            "D) It counts the number of punctuation marks in a document"
        ],
        "answer_index": 0,
        "explanation": "Coreference resolution links pronouns and aliases back to their canonical entity mentions (e.g., linking 'He founded SpaceX' back to 'Elon Musk'), enabling accurate cross-sentence relationship extraction."
    }
]


# =============================================================================
# 2. MODEL LOADING & FILE PARSING
# =============================================================================

@st.cache_resource(show_spinner=False)
def load_nlp_model(model_name: str = "en_core_web_sm"):
    """Loads spaCy model with caching and clear error handling."""
    if not SPACY_AVAILABLE:
        return None, "spaCy is not installed."
    try:
        nlp_model = spacy.load(model_name)
        return nlp_model, None
    except Exception as err:
        return None, str(err)


def extract_text_from_file(uploaded_file) -> Tuple[Optional[str], Optional[str]]:
    """Extracts raw text from .txt, .pdf, or .docx files."""
    if uploaded_file is None:
        return None, "No file uploaded."

    ext = os.path.splitext(uploaded_file.name)[1].lower()
    raw_bytes = uploaded_file.getvalue()

    try:
        if ext == ".txt":
            try:
                return raw_bytes.decode("utf-8").strip(), None
            except UnicodeDecodeError:
                return raw_bytes.decode("latin-1", errors="ignore").strip(), None

        elif ext == ".pdf":
            if not PYPDF_AVAILABLE:
                return None, "pypdf library not installed. Please run: pip install pypdf"
            pdf_reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
            pages = [page.extract_text() for page in pdf_reader.pages if page.extract_text()]
            return "\n\n".join(pages).strip(), None

        elif ext == ".docx":
            if not DOCX_AVAILABLE:
                return None, "python-docx library not installed. Please run: pip install python-docx"
            doc_obj = docx.Document(io.BytesIO(raw_bytes))
            paras = [p.text.strip() for p in doc_obj.paragraphs if p.text.strip()]
            return "\n\n".join(paras).strip(), None

        else:
            return None, f"Unsupported file extension '{ext}'. Please upload .txt, .pdf, or .docx."

    except Exception as e:
        return None, f"Error reading file: {str(e)}"


# =============================================================================
# 3. CORE RELATIONSHIP EXTRACTION ENGINE
# =============================================================================

ENTITY_COLORS = {
    "PERSON": "#8b5cf6",
    "ORG": "#3b82f6",
    "GPE": "#10b981",
    "LOC": "#059669",
    "DATE": "#f59e0b",
    "PRODUCT": "#ef4444",
    "CONCEPT": "#0284c7",
    "ENTITY": "#64748b"
}

def get_entity_color(ent_type: str) -> str:
    return ENTITY_COLORS.get(ent_type.upper(), "#64748b")


def clean_entity_text(text: str) -> str:
    """Removes leading/trailing whitespace and punctuation from entity names."""
    if not text:
        return ""
    return text.strip().strip(".,;:!?\"'`()[]{}").strip()


def resolve_entity_span(token, doc: Doc) -> Tuple[str, str, bool]:
    """Resolves a syntactic token to its named entity or noun chunk.
    Returns (clean_text, label, is_proper_or_named)."""
    # 1. Check doc.ents
    for ent in doc.ents:
        if ent.start <= token.i < ent.end:
            return clean_entity_text(ent.text), ent.label_, True

    # 2. Check noun chunks
    for chunk in doc.noun_chunks:
        if chunk.start <= token.i < chunk.end:
            clean_text = chunk.text
            if len(chunk) > 1 and (chunk[0].pos_ == "DET" or chunk[0].dep_ == "det"):
                clean_text = doc[chunk.start + 1: chunk.end].text.strip()
            if not clean_text:
                clean_text = chunk.text.strip()
            is_prop = any(t.pos_ == "PROPN" for t in chunk)
            label = "CONCEPT" if token.pos_ == "NOUN" else ("ORG" if token.pos_ == "PROPN" else "ENTITY")
            return clean_entity_text(clean_text), label, is_prop

    # 3. Expand compounds
    start_i = token.i
    end_i = token.i + 1
    for left in reversed(list(token.lefts)):
        if left.dep_ in ("compound", "amod", "nummod", "poss") or (left.pos_ in ("PROPN", "ADJ", "NOUN", "NUM") and left.dep_ != "punct"):
            start_i = min(start_i, left.i)
        else:
            break
    for right in token.rights:
        if right.dep_ in ("compound", "nummod") or (right.pos_ in ("PROPN", "NOUN", "NUM") and right.dep_ != "punct"):
            end_i = max(end_i, right.i + 1)
        else:
            break

    span_tokens = doc[start_i:end_i]
    span_text = span_tokens.text.strip()
    is_prop = any(t.pos_ == "PROPN" for t in span_tokens)
    label = "DATE" if (token.like_num or token.pos_ == "NUM") else ("ORG" if token.pos_ == "PROPN" else "CONCEPT")
    return clean_entity_text(span_text), label, is_prop


def is_date_or_num(text: str, ent_type: str) -> bool:
    """Checks if a span represents a date, number, or year fragment."""
    if ent_type in ("DATE", "TIME", "PERCENT", "MONEY", "QUANTITY", "CARDINAL", "ORDINAL"):
        return True
    clean_s = text.strip()
    if clean_s.isdigit() or re.match(r'^(18|19|20)\d{2}$', clean_s):
        return True
    return False


def normalize_relation(rel_str: str) -> str:
    """Normalizes verbal phrases into standard snake_case predicates."""
    s = rel_str.strip().lower()

    if "co-found" in s or "cofound" in s or "co_found" in s or s == "co-founded":
        return "co-founded"
    if "co-author" in s or "coauthor" in s or "co_author" in s:
        return "co-authored"
    if "co-creat" in s or "cocreat" in s or "co_creat" in s:
        return "co-created"
    if "co-direct" in s or "codirect" in s or "co_direct" in s:
        return "co-directed"
    if "co-own" in s or "coown" in s or "co_own" in s:
        return "co-owns"
    if "partner" in s:
        return "partner_with" if "with" in s else "partnered"
    if "succeed" in s:
        return "succeeded"
    if "born" in s or "bear" in s:
        return "born_in" if "in" in s else "born"
    if "died" in s or "die" in s:
        return "died_in" if "in" in s else "died"
    if "headquarter" in s:
        return "headquartered_in"
    if "found" in s:
        if "by" in s:
            return "founded"
        if "in" in s:
            return "founded_in"
        return "founded"
    if "acquir" in s:
        return "acquired"
    if "work" in s:
        return "works_for" if "for" in s else ("worked_at" if "at" in s else "works_at")
    if "stud" in s:
        return "studied_at" if "at" in s else "studied"
    if "graduat" in s:
        return "graduated_from"
    if "locat" in s:
        return "located_in"
    if "base" in s:
        return "based_in"
    if "marri" in s:
        return "married_to"
    if "lead" in s:
        return "leads"
    if "manufactur" in s:
        return "manufactures"
    if "produc" in s:
        return "produces"
    if "invent" in s:
        return "invented"
    if "creat" in s:
        return "created"
    if "develop" in s:
        return "developed_by" if "by" in s else "developed"
    if "discover" in s:
        return "discovered_in" if "in" in s else "discovered"
    if "win" in s or "won" in s:
        return "won"
    if "ceo" in s:
        return "CEO_of"
    if "president" in s:
        return "president_of"
    if "live" in s:
        return "lives_in"
    if "part" in s and "of" in s:
        return "part_of"
    if "member" in s and "of" in s:
        return "member_of"
    if "capital" in s and "of" in s:
        return "capital_of"
    if "own" in s:
        return "owns"
    if "run" in s:
        return "runs"
    if "join" in s:
        return "joined"

    stopwords = {"is", "was", "are", "were", "been", "has", "have", "had", "the", "a", "an", "also", "be"}
    words = [w for w in re.split(r'[\s_]+', s) if w and w not in stopwords]
    if not words:
        words = [w for w in re.split(r'[\s_]+', s) if w]
    return "_".join(words) if words else "relates_to"


def is_negated(token: Token) -> bool:
    """Checks if a token or its head has negative modifiers."""
    for child in token.children:
        if child.dep_ == "neg" or child.text.lower() in ("not", "no", "never", "neither", "nor", "none", "nobody", "nothing", "nowhere"):
            return True
        if child.dep_ == "det" and child.text.lower() in ("no", "neither"):
            return True
    if token.head and token.head != token:
        for child in token.head.children:
            if child.dep_ == "neg" and child.text.lower() in ("not", "no", "never"):
                return True
    return False


def get_conjuncts(token: Token) -> List[Token]:
    """Collects coordinate conjuncts (e.g. Steve Jobs, Steve Wozniak, and Ronald Wayne)."""
    conjs = [token]
    for child in token.children:
        if child.dep_ == "conj":
            conjs.extend(get_conjuncts(child))
    return conjs


def find_subject_for_verb(verb: Token) -> List[Token]:
    """Finds subject for a verb, checking children or parent verbs."""
    subjs = [c for c in verb.children if c.dep_ in ("nsubj", "nsubjpass", "csubj")]
    if subjs:
        return subjs
    curr = verb.head
    while curr and curr != verb and curr.pos_ in ("VERB", "AUX"):
        curr_subjs = [c for c in curr.children if c.dep_ in ("nsubj", "nsubjpass", "csubj")]
        if curr_subjs:
            return curr_subjs
        if curr.head == curr:
            break
        curr = curr.head
    return []


def preprocess_text_for_hyphenated_verbs(text: str) -> str:
    """Normalizes hyphenated prefixes (e.g., co-founded -> co_founded) so spaCy parses them cleanly as single verbs."""
    return re.sub(r'\b(co|re|pre|vice)-([a-zA-Z]+)\b', r'\1_\2', text)


def extract_relationships_from_sentence(sent: Span, doc: Doc) -> List[Dict[str, Any]]:
    """Extracts entity–relationship triples using dependency parsing."""
    triples: List[Dict[str, Any]] = []
    verbs = [tok for tok in sent if tok.pos_ in ("VERB", "AUX")]

    for verb in verbs:
        is_copula = verb.dep_ in ("ROOT", "ccomp", "conj") and any(c.dep_ in ("attr", "acomp", "prep") for c in verb.children)
        if verb.pos_ == "AUX" and not is_copula:
            continue

        if is_negated(verb):
            continue

        raw_subjects = find_subject_for_verb(verb)
        subjects: List[Token] = []
        for s in raw_subjects:
            if s.pos_ == "PRON" and s.text.lower() in ("which", "who", "that"):
                if verb.dep_ == "relcl" and verb.head:
                    subjects.append(verb.head)
                else:
                    subjects.append(s)
            else:
                subjects.append(s)

        is_passive = any(c.dep_ in ("auxpass", "nsubjpass") for c in verb.children) or (verb.dep_ == "nsubjpass")
        verb_str = verb.lemma_

        # A. Passive Voice: "Apple was founded by Steve Jobs..."
        if is_passive:
            passive_objs: List[Token] = []
            for s in subjects:
                for subj_conj in get_conjuncts(s):
                    if not is_negated(subj_conj):
                        passive_objs.append(subj_conj)

            agents: List[Token] = []
            for c in verb.children:
                if c.dep_ in ("agent", "prep") and c.text.lower() == "by":
                    for pobj in c.children:
                        if pobj.dep_ in ("pobj", "dobj") and not is_negated(pobj):
                            agents.extend(get_conjuncts(pobj))

            for c in verb.children:
                if c.dep_ == "conj" and c.pos_ == "PROPN" and not is_negated(c):
                    agents.append(c)

            for agent in agents:
                agent_text, agent_type, _ = resolve_entity_span(agent, doc)
                if not agent_text or agent_text.lower() in ("no company", "nobody", "none", "nothing", "no one"):
                    continue
                for p_obj in passive_objs:
                    obj_text, obj_type, _ = resolve_entity_span(p_obj, doc)
                    if obj_text and agent_text.lower() != obj_text.lower():
                        rel = normalize_relation(verb_str if verb.lemma_ != "be" else verb.text)
                        triples.append({
                            "subject": agent_text,
                            "subject_type": agent_type,
                            "relation": rel,
                            "object": obj_text,
                            "object_type": obj_type,
                            "sentence": sent.text
                        })

            preps = [c for c in verb.children if c.dep_ == "prep" and c.text.lower() != "by"]
            for prep in preps:
                prep_lemma = prep.text.lower()
                for pobj in prep.children:
                    if pobj.dep_ in ("pobj", "dobj") and not is_negated(pobj):
                        for p_conj in get_conjuncts(pobj):
                            p_text, p_type, p_is_prop = resolve_entity_span(p_conj, doc)
                            if is_date_or_num(p_text, p_type):
                                continue
                            # Avoid generic noun phrases in prepositional attachments
                            if not p_is_prop and p_type in ("CONCEPT", "ENTITY"):
                                continue
                            for p_obj in passive_objs:
                                obj_text, obj_type, _ = resolve_entity_span(p_obj, doc)
                                if p_text and obj_text and p_text.lower() != obj_text.lower():
                                    rel = normalize_relation(f"{verb_str} {prep_lemma}")
                                    triples.append({
                                        "subject": obj_text,
                                        "subject_type": obj_type,
                                        "relation": rel,
                                        "object": p_text,
                                        "object_type": p_type,
                                        "sentence": sent.text
                                    })
            continue

        # B. Active Voice: "Elon Musk founded SpaceX" / "NASA partnered with SpaceX"
        direct_objects: List[Token] = []
        for c in verb.children:
            if c.dep_ in ("dobj", "obj") and not is_negated(c):
                direct_objects.extend(get_conjuncts(c))

        attr_objects: List[Tuple[Token, Optional[Token], str]] = []
        for c in verb.children:
            if c.dep_ in ("attr", "acomp"):
                for p in c.children:
                    if p.dep_ == "prep" and p.text.lower() == "of":
                        for pobj in p.children:
                            if pobj.dep_ in ("pobj", "dobj") and not is_negated(pobj):
                                attr_objects.append((c, pobj, f"{c.text}_of"))
                if not attr_objects and not is_negated(c):
                    attr_objects.append((c, None, f"is_{c.lemma_}"))

        for s in subjects:
            for s_conj in get_conjuncts(s):
                if is_negated(s_conj):
                    continue
                s_text, s_type, _ = resolve_entity_span(s_conj, doc)
                if not s_text:
                    continue

                for o in direct_objects:
                    o_text, o_type, _ = resolve_entity_span(o, doc)
                    if o_text and s_text.lower() != o_text.lower():
                        if is_date_or_num(o_text, o_type):
                            continue
                        rel = normalize_relation(verb_str)
                        triples.append({
                            "subject": s_text,
                            "subject_type": s_type,
                            "relation": rel,
                            "object": o_text,
                            "object_type": o_type,
                            "sentence": sent.text
                        })

                for c in verb.children:
                    if c.dep_ == "prep":
                        prep_lemma = c.text.lower()
                        # Special clean handling for "as <Role> of <Org>" (e.g. "as CEO of Apple")
                        if prep_lemma == "as":
                            for pobj in c.children:
                                if pobj.dep_ in ("pobj", "dobj") and not is_negated(pobj):
                                    of_preps = [p for p in pobj.children if p.dep_ == "prep" and p.text.lower() == "of"]
                                    for of_p in of_preps:
                                        for org_obj in of_p.children:
                                            if org_obj.dep_ in ("pobj", "dobj") and not is_negated(org_obj):
                                                for org_conj in get_conjuncts(org_obj):
                                                    org_text, org_type, _ = resolve_entity_span(org_conj, doc)
                                                    if org_text and s_text.lower() != org_text.lower():
                                                        role_name = pobj.text.strip()
                                                        triples.append({
                                                            "subject": s_text,
                                                            "subject_type": s_type,
                                                            "relation": normalize_relation(f"{role_name}_of"),
                                                            "object": org_text,
                                                            "object_type": org_type,
                                                            "sentence": sent.text
                                                        })
                            continue

                        # When a verb has a direct object, a "with" preposition represents comitative accompaniment,
                        # not the entity being acted upon (e.g., "Steve Wozniak co-founded Apple with Steve Jobs").
                        if prep_lemma == "with" and direct_objects:
                            continue

                        for pobj in c.children:
                            if pobj.dep_ in ("pobj", "dobj") and not is_negated(pobj):
                                for p_conj in get_conjuncts(pobj):
                                    p_text, p_type, p_is_prop = resolve_entity_span(p_conj, doc)
                                    if is_date_or_num(p_text, p_type):
                                        continue
                                    # Avoid unnecessary generic-noun prepositional attachments
                                    if not p_is_prop and p_type in ("CONCEPT", "ENTITY"):
                                        continue
                                    if p_text and s_text.lower() != p_text.lower():
                                        rel = normalize_relation(f"{verb_str} {prep_lemma}")
                                        triples.append({
                                            "subject": s_text,
                                            "subject_type": s_type,
                                            "relation": rel,
                                            "object": p_text,
                                            "object_type": p_type,
                                            "sentence": sent.text
                                        })

                for attr_tok, pobj_tok, rel_lbl in attr_objects:
                    if pobj_tok:
                        for p_conj in get_conjuncts(pobj_tok):
                            p_text, p_type, _ = resolve_entity_span(p_conj, doc)
                            if is_date_or_num(p_text, p_type):
                                continue
                            if p_text and s_text.lower() != p_text.lower():
                                triples.append({
                                    "subject": s_text,
                                    "subject_type": s_type,
                                    "relation": normalize_relation(rel_lbl),
                                    "object": p_text,
                                    "object_type": p_type,
                                    "sentence": sent.text
                                })

    # C. Appositive Locations: "Cupertino, California"
    for tok in sent:
        if tok.dep_ == "appos":
            head_tok = tok.head
            head_text, head_type, _ = resolve_entity_span(head_tok, doc)
            appos_text, appos_type, _ = resolve_entity_span(tok, doc)
            if head_text and appos_text and head_text.lower() != appos_text.lower():
                if head_type in ("GPE", "LOC", "ORG") or appos_type in ("GPE", "LOC"):
                    triples.append({
                        "subject": head_text,
                        "subject_type": head_type,
                        "relation": "located_in",
                        "object": appos_text,
                        "object_type": appos_type,
                        "sentence": sent.text
                    })

    # Deduplicate triples and clean punctuation
    seen = set()
    unique_triples = []
    for t in triples:
        subj = clean_entity_text(t["subject"])
        obj = clean_entity_text(t["object"])
        rel = t["relation"]
        if subj and obj:
            key = (subj.lower(), rel.lower(), obj.lower())
            if key not in seen:
                seen.add(key)
                unique_triples.append({
                    "subject": subj,
                    "subject_type": t.get("subject_type", "ENTITY"),
                    "relation": rel,
                    "object": obj,
                    "object_type": t.get("object_type", "ENTITY"),
                    "sentence": t.get("sentence", "")
                })

    return unique_triples


def extract_entities_list(doc: Doc) -> List[Dict[str, str]]:
    """Returns detected named entities with cleaned punctuation."""
    seen = set()
    entities = []
    for ent in doc.ents:
        clean_name = clean_entity_text(ent.text)
        key = (clean_name, ent.label_)
        if key not in seen and clean_name and not is_date_or_num(clean_name, ent.label_):
            seen.add(key)
            entities.append({"Entity": clean_name, "Type": ent.label_})
    return entities


# =============================================================================
# 4. KNOWLEDGE GRAPH RENDERER
# =============================================================================

def render_knowledge_graph(triples: List[Dict[str, Any]]):
    """Renders a clean, clear interactive Knowledge Graph."""
    if not triples:
        return

    G = nx.DiGraph()
    for t in triples:
        subj = t["subject"]
        obj = t["object"]
        rel = t["relation"]
        subj_type = t.get("subject_type", "ENTITY")
        obj_type = t.get("object_type", "ENTITY")

        if not G.has_node(subj):
            G.add_node(subj, color=get_entity_color(subj_type), title=f"{subj} ({subj_type})")
        if not G.has_node(obj):
            G.add_node(obj, color=get_entity_color(obj_type), title=f"{obj} ({obj_type})")

        G.add_edge(subj, obj, label=rel, title=rel)

    if PYVIS_AVAILABLE:
        try:
            net = Network(height="420px", width="100%", directed=True, bgcolor="#ffffff", font_color="#0f172a")
            for node_id, data in G.nodes(data=True):
                net.add_node(
                    node_id,
                    label=node_id,
                    title=data.get("title", node_id),
                    color=data.get("color", "#3b82f6"),
                    size=20,
                    shape="dot",
                    font={"size": 13, "color": "#0f172a", "face": "Inter, sans-serif"}
                )
            for u, v, data in G.edges(data=True):
                net.add_edge(
                    u, v,
                    label=data.get("label", ""),
                    color="#94a3b8",
                    arrows="to",
                    font={"size": 11, "color": "#334155", "align": "middle", "background": "#f8fafc"}
                )
            net.barnes_hut(gravity=-2000, central_gravity=0.3, spring_length=130, spring_strength=0.05)
            html_code = net.generate_html()
            components.html(html_code, height=440, scrolling=False)
            return
        except Exception:
            pass

    # Fallback to Matplotlib
    if MATPLOTLIB_AVAILABLE:
        fig, ax = plt.subplots(figsize=(8, 4.5), facecolor="#ffffff")
        pos = nx.spring_layout(G, seed=42)
        node_colors = [d.get("color", "#3b82f6") for _, d in G.nodes(data=True)]
        nx.draw_networkx_nodes(G, pos, node_color=node_colors, node_size=1200, alpha=0.9, ax=ax)
        nx.draw_networkx_labels(G, pos, font_size=8, font_family="sans-serif", font_weight="bold", ax=ax)
        nx.draw_networkx_edges(G, pos, edge_color="#94a3b8", arrows=True, arrowsize=14, ax=ax)
        edge_labels = {(u, v): d.get("label", "") for u, v, d in G.edges(data=True)}
        nx.draw_networkx_edge_labels(G, pos, edge_labels=edge_labels, font_size=7, ax=ax)
        plt.axis("off")
        st.pyplot(fig)
        plt.close(fig)


# =============================================================================
# 5. LAB REPORT PDF EXPORTER
# =============================================================================

def get_unicode_font_paths() -> Dict[str, Optional[str]]:
    """Locates TrueType fonts supporting full Unicode character sets (DejaVu Sans or system TTF)."""
    # Check 1: Local project fonts directory
    local_fonts_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
    dejavu_reg = os.path.join(local_fonts_dir, "DejaVuSans.ttf")
    dejavu_bold = os.path.join(local_fonts_dir, "DejaVuSans-Bold.ttf")
    dejavu_ital = os.path.join(local_fonts_dir, "DejaVuSans-Oblique.ttf")
    dejavu_bi = os.path.join(local_fonts_dir, "DejaVuSans-BoldOblique.ttf")
    if os.path.exists(dejavu_reg):
        return {
            "regular": dejavu_reg,
            "bold": dejavu_bold if os.path.exists(dejavu_bold) else dejavu_reg,
            "italic": dejavu_ital if os.path.exists(dejavu_ital) else dejavu_reg,
            "bold_italic": dejavu_bi if os.path.exists(dejavu_bi) else None,
            "family": "DejaVu"
        }

    # Check 2: Matplotlib bundled fonts
    if MATPLOTLIB_AVAILABLE:
        try:
            import matplotlib
            mpl_font_dir = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
            mpl_reg = os.path.join(mpl_font_dir, "DejaVuSans.ttf")
            mpl_bold = os.path.join(mpl_font_dir, "DejaVuSans-Bold.ttf")
            mpl_ital = os.path.join(mpl_font_dir, "DejaVuSans-Oblique.ttf")
            mpl_bi = os.path.join(mpl_font_dir, "DejaVuSans-BoldOblique.ttf")
            if os.path.exists(mpl_reg):
                return {
                    "regular": mpl_reg,
                    "bold": mpl_bold if os.path.exists(mpl_bold) else mpl_reg,
                    "italic": mpl_ital if os.path.exists(mpl_ital) else mpl_reg,
                    "bold_italic": mpl_bi if os.path.exists(mpl_bi) else None,
                    "family": "DejaVu"
                }
        except Exception:
            pass

        try:
            import matplotlib.font_manager as fm
            reg = fm.findfont(fm.FontProperties(family="DejaVu Sans", weight="normal", style="normal"))
            bold = fm.findfont(fm.FontProperties(family="DejaVu Sans", weight="bold", style="normal"))
            ital = fm.findfont(fm.FontProperties(family="DejaVu Sans", weight="normal", style="italic"))
            if reg and os.path.exists(reg) and "dejavu" in reg.lower():
                return {
                    "regular": reg,
                    "bold": bold if bold and os.path.exists(bold) else reg,
                    "italic": ital if ital and os.path.exists(ital) else reg,
                    "bold_italic": None,
                    "family": "DejaVu"
                }
        except Exception:
            pass

    # Check 3: Standard Windows / Linux / macOS paths
    system_candidates = [
        {
            "family": "Arial",
            "regular": r"C:\Windows\Fonts\arial.ttf",
            "bold": r"C:\Windows\Fonts\arialbd.ttf",
            "italic": r"C:\Windows\Fonts\ariali.ttf",
            "bold_italic": r"C:\Windows\Fonts\arialbi.ttf"
        },
        {
            "family": "SegoeUI",
            "regular": r"C:\Windows\Fonts\segoeui.ttf",
            "bold": r"C:\Windows\Fonts\segoeuib.ttf",
            "italic": r"C:\Windows\Fonts\segoeuii.ttf",
            "bold_italic": r"C:\Windows\Fonts\segoeuiz.ttf"
        },
        {
            "family": "DejaVu",
            "regular": "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "bold": "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "italic": "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
            "bold_italic": "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"
        }
    ]
    for cand in system_candidates:
        if os.path.exists(cand["regular"]):
            return {
                "regular": cand["regular"],
                "bold": cand["bold"] if os.path.exists(cand["bold"]) else cand["regular"],
                "italic": cand["italic"] if os.path.exists(cand["italic"]) else cand["regular"],
                "bold_italic": cand["bold_italic"] if os.path.exists(cand["bold_italic"]) else None,
                "family": cand["family"]
            }

    return {"family": "Helvetica", "regular": None, "bold": None, "italic": None, "bold_italic": None}


def sanitize_pdf_text(text: Any, is_unicode: bool = True) -> str:
    """Sanitizes text strings for PDF generation, preserving full Unicode glyphs when supported."""
    if text is None:
        return ""
    s = str(text)
    if is_unicode:
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', s)

    replacements = {
        "\u2013": "-",
        "\u2014": "--",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2022": "*",
        "\u2026": "...",
        "\u00a0": " ",
        "\u2212": "-",
        "\u2192": "->"
    }
    for k, v in replacements.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


if FPDF_AVAILABLE:
    class LabReportPDF(FPDF):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.font_family_name = "Helvetica"
            self._setup_unicode_fonts()

        def _setup_unicode_fonts(self):
            font_info = get_unicode_font_paths()
            fam = font_info.get("family", "Helvetica")
            reg_path = font_info.get("regular")
            if reg_path and os.path.exists(reg_path):
                try:
                    self.add_font(fam, "", reg_path)
                    bold_path = font_info.get("bold")
                    if bold_path and os.path.exists(bold_path):
                        self.add_font(fam, "B", bold_path)
                    ital_path = font_info.get("italic")
                    if ital_path and os.path.exists(ital_path):
                        self.add_font(fam, "I", ital_path)
                    bi_path = font_info.get("bold_italic")
                    if bi_path and os.path.exists(bi_path):
                        self.add_font(fam, "BI", bi_path)
                    self.font_family_name = fam
                except Exception:
                    self.font_family_name = "Helvetica"
            else:
                self.font_family_name = "Helvetica"

        def footer(self):
            self.set_y(-15)
            self.set_font(self.font_family_name, "I", 8)
            self.set_text_color(128, 128, 128)
            self.cell(0, 10, f"Page {self.page_no()}/{{nb}} | Virtual Laboratory Report", align="C")


def generate_pdf_report(student_name: str, student_id: str, date_str: str,
                        entities_list: List[Dict[str, str]], triples_list: List[Dict[str, Any]],
                        quiz_score: int, quiz_total: int, student_notes: str) -> bytes:
    """Compiles relationship extraction experiment records into a formatted PDF report."""
    if not FPDF_AVAILABLE:
        return b""

    temp_img_path = None
    try:
        # Generate Knowledge Graph temporary image if triples exist
        if triples_list and NETWORKX_AVAILABLE and MATPLOTLIB_AVAILABLE:
            G = nx.DiGraph()
            for t in triples_list:
                subj = t["subject"]
                obj = t["object"]
                rel = t["relation"]
                subj_type = t.get("subject_type", "ENTITY")
                obj_type = t.get("object_type", "ENTITY")

                if not G.has_node(subj):
                    G.add_node(subj, color=get_entity_color(subj_type))
                if not G.has_node(obj):
                    G.add_node(obj, color=get_entity_color(obj_type))

                G.add_edge(subj, obj, label=rel)

            fig, ax = plt.subplots(figsize=(8, 3.6), facecolor="#ffffff")
            pos = nx.spring_layout(G, seed=42)
            node_colors = [d.get("color", "#3b82f6") for _, d in G.nodes(data=True)]
            nx.draw_networkx_nodes(G, pos, node_color=node_colors, node_size=1100, alpha=0.9, ax=ax)
            nx.draw_networkx_labels(G, pos, font_size=8, font_family="sans-serif", font_weight="bold", ax=ax)
            nx.draw_networkx_edges(G, pos, edge_color="#94a3b8", arrows=True, arrowsize=14, ax=ax)
            edge_labels = {(u, v): d.get("label", "") for u, v, d in G.edges(data=True)}
            nx.draw_networkx_edge_labels(G, pos, edge_labels=edge_labels, font_size=7, ax=ax)
            plt.axis("off")

            temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            temp_img_path = temp_file.name
            temp_file.close()

            plt.savefig(temp_img_path, format="png", dpi=200, bbox_inches="tight")
            plt.close(fig)

        pdf = LabReportPDF()
        pdf.alias_nb_pages()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        font_fam = pdf.font_family_name
        is_unicode = (font_fam != "Helvetica")

        # Document Header
        pdf.set_text_color(15, 23, 42)
        pdf.set_font(font_fam, "B", 14)
        pdf.cell(0, 7, "Virtual Laboratory Report", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(font_fam, "B", 11)
        pdf.set_text_color(30, 58, 138)
        pdf.cell(0, 6, sanitize_pdf_text(f"Experiment: {EXPERIMENT_CONFIG['title']}", is_unicode), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

        # Student & Session Info Box
        pdf.set_fill_color(241, 245, 249)
        pdf.set_draw_color(203, 213, 225)
        pdf.rect(10, 25, 190, 20, "FD")

        pdf.set_xy(14, 27)
        pdf.set_font(font_fam, "B", 8.5)
        pdf.set_text_color(71, 85, 105)
        pdf.cell(38, 4.5, "Student Name:", new_x="RIGHT", new_y="TOP")
        pdf.set_font(font_fam, "", 8.5)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(57, 4.5, sanitize_pdf_text(student_name or "N/A", is_unicode), new_x="RIGHT", new_y="TOP")

        pdf.set_font(font_fam, "B", 8.5)
        pdf.set_text_color(71, 85, 105)
        pdf.cell(35, 4.5, "Student ID / Roll:", new_x="RIGHT", new_y="TOP")
        pdf.set_font(font_fam, "", 8.5)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(50, 4.5, sanitize_pdf_text(student_id or "N/A", is_unicode), new_x="LMARGIN", new_y="NEXT")

        pdf.set_xy(14, 34)
        pdf.set_font(font_fam, "B", 8.5)
        pdf.set_text_color(71, 85, 105)
        pdf.cell(38, 4.5, "Experiment Date:", new_x="RIGHT", new_y="TOP")
        pdf.set_font(font_fam, "", 8.5)
        pdf.set_text_color(15, 23, 42)
        pdf.cell(57, 4.5, sanitize_pdf_text(date_str or datetime.now().strftime("%Y-%m-%d"), is_unicode), new_x="RIGHT", new_y="TOP")

        pdf.set_font(font_fam, "B", 8.5)
        pdf.set_text_color(71, 85, 105)
        pdf.cell(35, 4.5, "Quiz Evaluation:", new_x="RIGHT", new_y="TOP")
        pdf.set_font(font_fam, "B", 8.5)
        if quiz_score >= max(1, quiz_total // 2):
            pdf.set_text_color(16, 185, 129)
        else:
            pdf.set_text_color(239, 68, 68)
        pdf.cell(50, 4.5, f"{quiz_score} / {quiz_total} ({int((quiz_score/quiz_total)*100 if quiz_total else 0)}%)", new_x="LMARGIN", new_y="NEXT")

        pdf.ln(10)

        # 1. Learning Objectives
        pdf.set_font(font_fam, "B", 10.5)
        pdf.set_text_color(30, 58, 138)
        pdf.cell(0, 6, "1. Learning Objectives", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(font_fam, "", 8.5)
        pdf.set_text_color(51, 65, 85)
        for obj in EXPERIMENT_CONFIG["objectives"]:
            clean_obj = sanitize_pdf_text(str(obj).replace("$", "").replace("\\", ""), is_unicode)
            pdf.cell(4, 4.2, "-", new_x="RIGHT", new_y="TOP")
            pdf.cell(0, 4.2, f" {clean_obj}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

        # 2. Extracted Entities Summary
        pdf.set_font(font_fam, "B", 10.5)
        pdf.set_text_color(30, 58, 138)
        pdf.cell(0, 6, "2. Extracted Named Entities", new_x="LMARGIN", new_y="NEXT")

        if entities_list:
            ent_str = ", ".join([f"{e['Entity']} [{e['Type']}]" for e in entities_list])
            pdf.set_font(font_fam, "", 8.5)
            pdf.set_text_color(51, 65, 85)
            pdf.multi_cell(0, 4.5, sanitize_pdf_text(ent_str, is_unicode))
        else:
            pdf.set_font(font_fam, "I", 8.5)
            pdf.set_text_color(100, 116, 139)
            pdf.cell(0, 5, "No named entities recorded during this session.", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

        # 3. Extracted Relationship Triples Table
        pdf.set_font(font_fam, "B", 10.5)
        pdf.set_text_color(30, 58, 138)
        pdf.cell(0, 6, "3. Extracted Semantic Triples (Subject – Predicate – Object)", new_x="LMARGIN", new_y="NEXT")

        if not triples_list:
            pdf.set_font(font_fam, "I", 8.5)
            pdf.set_text_color(100, 116, 139)
            pdf.cell(0, 5, "No relationship triples extracted during this session.", new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_fill_color(37, 99, 235)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font(font_fam, "B", 8)

            col_w_subj = 60
            col_w_rel = 55
            col_w_obj = 75

            pdf.cell(col_w_subj, 5.5, "Subject (e1)", 1, new_x="RIGHT", new_y="TOP", align="C", fill=True)
            pdf.cell(col_w_rel, 5.5, "Relationship (r)", 1, new_x="RIGHT", new_y="TOP", align="C", fill=True)
            pdf.cell(col_w_obj, 5.5, "Object (e2)", 1, new_x="LMARGIN", new_y="NEXT", align="C", fill=True)

            pdf.set_fill_color(248, 250, 252)
            pdf.set_text_color(30, 41, 59)
            pdf.set_font(font_fam, "", 7.5)
            fill = False

            for t in triples_list:
                pdf.cell(col_w_subj, 4.8, sanitize_pdf_text(str(t.get("subject", ""))[:32], is_unicode), 1, new_x="RIGHT", new_y="TOP", align="L", fill=fill)
                pdf.cell(col_w_rel, 4.8, sanitize_pdf_text(str(t.get("relation", ""))[:28], is_unicode), 1, new_x="RIGHT", new_y="TOP", align="C", fill=fill)
                pdf.cell(col_w_obj, 4.8, sanitize_pdf_text(str(t.get("object", ""))[:42], is_unicode), 1, new_x="LMARGIN", new_y="NEXT", align="L", fill=fill)
                fill = not fill
        pdf.ln(4)

        # 4. Knowledge Graph (only included if triples exist)
        sec_num = 4
        if triples_list and temp_img_path and os.path.exists(temp_img_path):
            if pdf.get_y() + 75 > (pdf.h - 15):
                pdf.add_page()
            pdf.set_font(font_fam, "B", 10.5)
            pdf.set_text_color(30, 58, 138)
            pdf.cell(0, 6, f"{sec_num}. Knowledge Graph", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
            img_w = 150
            x_pos = (pdf.w - img_w) / 2
            pdf.image(temp_img_path, x=x_pos, w=img_w)
            pdf.ln(4)
            sec_num += 1

        if pdf.get_y() + 35 > (pdf.h - 15):
            pdf.add_page()

        # Observations & Analysis
        pdf.set_font(font_fam, "B", 10.5)
        pdf.set_text_color(30, 58, 138)
        pdf.cell(0, 6, f"{sec_num}. Observations & Analysis", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(font_fam, "", 8.5)
        pdf.set_text_color(51, 65, 85)
        notes_text = student_notes.strip() if student_notes.strip() else (
            "The relationship extraction model successfully recognized named entities and mapped semantic relationships "
            "into structured knowledge graph triples, handling active/passive syntax and filtering negative assertions."
        )
        pdf.multi_cell(0, 4.5, sanitize_pdf_text(notes_text, is_unicode))
        pdf.ln(6)

        # Sign-off line
        if pdf.get_y() + 20 > (pdf.h - 15):
            pdf.add_page()
        pdf.set_draw_color(180, 180, 180)
        pdf.line(130, pdf.get_y() + 10, 190, pdf.get_y() + 10)
        pdf.set_xy(130, pdf.get_y() + 11)
        pdf.set_font(font_fam, "I", 7.5)
        pdf.set_text_color(100, 100, 100)
        pdf.cell(60, 4, "Instructor / Student Signature", align="C")

        return bytes(pdf.output())

    finally:
        if temp_img_path and os.path.exists(temp_img_path):
            try:
                os.remove(temp_img_path)
            except Exception:
                pass


# =============================================================================
# 6. SECTION RENDERERS: THEORY, SIMULATION, QUIZ, REPORT
# =============================================================================

def render_theory_section():
    """Renders Section 1: Theory, Background, Objectives, and Procedure."""
    st.header("Theoretical Framework & Background")
    st.markdown(THEORY_CONTENT["background"])

    st.subheader("Learning Objectives")
    for i, obj in enumerate(EXPERIMENT_CONFIG["objectives"]):
        st.write(f"- **Goal {i+1}**: {obj}")

    st.divider()
    st.subheader("Experimental Procedure")
    for step in THEORY_CONTENT["procedure"]:
        st.write(f"- {step}")

    st.divider()
    with st.expander("Key Terminology & Concept Reference"):
        var_df = pd.DataFrame(
            list(THEORY_CONTENT["key_terms"].items()),
            columns=["Term / Concept", "Definition & Role"]
        )
        st.table(var_df)


def render_simulation_section(nlp):
    """Renders Section 2: Interactive Relationship Extraction Sandbox & Graph."""
    st.header("Relationship Extraction Simulation")
    st.write("Select or enter input text below, then click **Extract Relationships** to run the pipeline.")

    # Input Mode Selection
    input_mode = st.radio(
        "Select Input Mode:",
        options=["Sample Example", "Enter Text", "Upload File"],
        horizontal=True
    )

    if input_mode == "Sample Example":
        selected_sample = st.selectbox(
            "Choose a sample example:",
            options=list(SAMPLE_TEXTS.keys()),
            index=0
        )
        if selected_sample != st.session_state.get("prev_sample_key") or st.session_state.get("prev_input_mode") != "Sample Example":
            st.session_state["input_text"] = SAMPLE_TEXTS[selected_sample]
            st.session_state["prev_sample_key"] = selected_sample

    elif input_mode == "Enter Text":
        if st.session_state.get("prev_input_mode") != "Enter Text":
            st.session_state["input_text"] = ""

    elif input_mode == "Upload File":
        uploaded_file = st.file_uploader(
            "Upload File (.txt, .pdf, .docx):",
            type=["txt", "pdf", "docx"]
        )
        if uploaded_file is not None:
            if uploaded_file.name != st.session_state.get("prev_uploaded_file"):
                extracted_txt, err = extract_text_from_file(uploaded_file)
                if err:
                    st.error(err)
                else:
                    st.session_state["input_text"] = extracted_txt
                    st.session_state["prev_uploaded_file"] = uploaded_file.name
            st.success("File loaded successfully.")
        else:
            if st.session_state.get("prev_input_mode") != "Upload File":
                st.session_state["input_text"] = ""

    st.session_state["prev_input_mode"] = input_mode

    # Input Text Area
    user_text = st.text_area(
        label="Input Text for Extraction:",
        value=st.session_state.get("input_text", ""),
        height=120,
        placeholder="Enter or paste text to extract entity relationships..."
    )
    st.session_state["input_text"] = user_text

    # Extract Relationships Button
    col_btn, _ = st.columns([2, 5])
    with col_btn:
        extract_clicked = st.button("Extract Relationships", type="primary", use_container_width=True)

    if extract_clicked:
        clean_text = user_text.strip()
        if not clean_text:
            st.warning("Please enter or select text before extracting relationships.")
            st.session_state["results"] = None
        else:
            with st.spinner("Extracting entities and relationships..."):
                preprocessed = preprocess_text_for_hyphenated_verbs(clean_text)
                doc = nlp(preprocessed)
                entities = extract_entities_list(doc)
                triples = []
                for sent in doc.sents:
                    triples.extend(extract_relationships_from_sentence(sent, doc))

                st.session_state["results"] = {
                    "text": clean_text,
                    "entities": entities,
                    "triples": triples
                }
                st.toast("Extraction completed successfully!")

    # Display Results (persisted in session_state)
    if st.session_state.get("results") is not None:
        st.divider()
        st.subheader("Extraction Results")

        res = st.session_state["results"]
        entities = res["entities"]
        triples = res["triples"]

        # 1. Entities Detected
        st.markdown("#### 1. Named Entities Detected")
        if entities:
            ent_df = pd.DataFrame(entities)
            col_e1, col_e2 = st.columns([3, 1])
            with col_e1:
                st.dataframe(ent_df, hide_index=True, use_container_width=True)
            with col_e2:
                st.metric("Total Entities", len(entities))
        else:
            st.info("No named entities detected.")

        # 2. Relationship Triples
        st.markdown("#### 2. Extracted Relationship Triples")
        if triples:
            triples_df = pd.DataFrame([
                {
                    "Subject": t["subject"],
                    "Relationship": t["relation"],
                    "Object": t["object"]
                }
                for t in triples
            ])
            st.dataframe(
                triples_df,
                use_container_width=True,
                height=280,
                hide_index=True,
                column_config={
                    "Subject": st.column_config.TextColumn(
                        "Subject",
                        width="medium",
                    ),
                    "Relationship": st.column_config.TextColumn(
                        "Relationship",
                        width="medium",
                    ),
                    "Object": st.column_config.TextColumn(
                        "Object",
                        width="medium",
                    ),
                }
            )

            # 3. Knowledge Graph
            st.markdown("#### 3. Knowledge Graph Visualization")
            render_knowledge_graph(triples)

            # 4. Downloads
            col_csv, col_json, _ = st.columns([1.5, 1.5, 3])
            with col_csv:
                csv_bytes = triples_df.to_csv(index=False).encode('utf-8')
                st.download_button(
                    label="Download CSV",
                    data=csv_bytes,
                    file_name="relationship_triples.csv",
                    mime="text/csv",
                    use_container_width=True
                )

            with col_json:
                json_bytes = json.dumps([
                    {
                        "Subject": t["subject"],
                        "Relationship": t["relation"],
                        "Object": t["object"]
                    }
                    for t in triples
                ], indent=2).encode('utf-8')
                st.download_button(
                    label="Download JSON",
                    data=json_bytes,
                    file_name="relationship_triples.json",
                    mime="application/json",
                    use_container_width=True
                )
        else:
            st.info("No relationship triples could be extracted from the provided text.")


def render_quiz_section():
    """Renders Section 3: Assessment Quiz with Self-Grading and Feedback."""
    st.header("Concept Assessment Quiz")
    st.write("Answer the conceptual questions below to test your understanding of Relationship Extraction and Knowledge Graphs.")

    if "current_quiz_questions" not in st.session_state or len(st.session_state.get("current_quiz_questions", [])) != 5:
        st.session_state["current_quiz_questions"] = random.sample(QUIZ_QUESTIONS, 5)

    current_questions = st.session_state["current_quiz_questions"]

    with st.form("lab_quiz_form"):
        user_responses = {}
        for idx, q in enumerate(current_questions, start=1):
            st.subheader(f"Question {idx}")
            st.write(q["question"])
            selected = st.radio(
                label=f"Options for Question {idx}:",
                options=q["options"],
                index=st.session_state["quiz_answers"].get(q["id"], 0),
                key=f"quiz_radio_{q['id']}",
                label_visibility="collapsed"
            )
            user_responses[q["id"]] = q["options"].index(selected)

        submitted = st.form_submit_button("Submit Quiz for Grading", type="primary")

    if submitted:
        score = 0
        st.session_state["quiz_answers"] = user_responses
        st.session_state["quiz_submitted"] = True

        st.divider()
        st.subheader("Evaluation Results and Feedback")
        for idx, q in enumerate(current_questions, start=1):
            user_ans = user_responses.get(q["id"])
            correct_ans = q["answer_index"]
            if user_ans == correct_ans:
                score += 1
                st.success(f"**Question {idx}: Correct!**\n\n_{q['explanation']}_")
            else:
                st.error(f"**Question {idx}: Incorrect.** (Your answer: {q['options'][user_ans]})\n\n"
                         f"**Correct Answer:** {q['options'][correct_ans]}\n\n"
                         f"**Reasoning:** _{q['explanation']}_")

        st.session_state["quiz_score"] = score
        perc = (score / len(current_questions)) * 100
        st.info(f"Final Score: **{score} / {len(current_questions)}** ({perc:.0f}%)")

    elif st.session_state.get("quiz_submitted", False):
        score = st.session_state.get("quiz_score", 0)
        perc = (score / len(current_questions)) * 100
        st.success(f"Quiz already submitted. Current score: **{score} / {len(current_questions)}** ({perc:.0f}%)")

    if st.session_state.get("quiz_submitted", False):
        if st.button("Retake Quiz (New Attempt)"):
            st.session_state["current_quiz_questions"] = random.sample(QUIZ_QUESTIONS, 5)
            st.session_state["quiz_answers"] = {}
            st.session_state["quiz_submitted"] = False
            st.session_state["quiz_score"] = 0
            st.rerun()


def render_report_section():
    """Renders Section 4: Dynamic Lab Report Generator with PDF Export."""
    st.header("Report Generation")
    st.write("Compile your student details, extracted relationship triples, and quiz evaluation into an official PDF report.")

    col1, col2, col3 = st.columns(3)
    with col1:
        student_name = st.text_input("Student Name", value=st.session_state["student_info"].get("name", "Student Name"))
    with col2:
        student_id = st.text_input("Student Roll / ID", value=st.session_state["student_info"].get("id", "EXP-001"))
    with col3:
        lab_date = st.date_input("Experiment Date", value=datetime.now())

    st.session_state["student_info"]["name"] = student_name
    st.session_state["student_info"]["id"] = student_id
    st.session_state["student_info"]["date"] = str(lab_date)

    st.subheader("Discussion & Observations")
    student_notes = st.text_area(
        "Enter your interpretation of results, observations, and conclusions:",
        value=st.session_state.get("student_notes", (
            "The relationship extraction model successfully recognized named entities and mapped semantic relationships "
            "into structured knowledge graph triples, handling active/passive syntax and filtering negative assertions."
        )),
        height=110
    )
    st.session_state["student_notes"] = student_notes

    results = st.session_state.get("results")
    entities = results.get("entities", []) if results else []
    triples = results.get("triples", []) if results else []

    quiz_total = len(st.session_state.get("current_quiz_questions", [])) or 5

    st.divider()
    st.subheader("Report Summary Preview")
    st.write(f"**Experiment:** {EXPERIMENT_CONFIG['title']}")
    st.write(f"**Student:** {student_name} | **ID:** {student_id} | **Date:** {lab_date}")
    st.write(f"**Quiz Score:** {st.session_state.get('quiz_score', 0)} / {quiz_total}")

    if triples:
        st.write(f"**Extracted Triples ({len(triples)}):**")
        triples_df = pd.DataFrame([
            {"Subject": t["subject"], "Relationship": t["relation"], "Object": t["object"]}
            for t in triples
        ])
        st.dataframe(
            triples_df,
            use_container_width=True,
            height=280,
            hide_index=True,
            column_config={
                "Subject": st.column_config.TextColumn(
                    "Subject",
                    width="medium",
                ),
                "Relationship": st.column_config.TextColumn(
                    "Relationship",
                    width="medium",
                ),
                "Object": st.column_config.TextColumn(
                    "Object",
                    width="medium",
                ),
            }
        )
    else:
        st.info("Note: You have not extracted any relationships in the Simulation tab yet. Your report will indicate 0 triples.")

    # Generate PDF bytes and write file to disk
    pdf_bytes = generate_pdf_report(
        student_name=student_name,
        student_id=student_id,
        date_str=str(lab_date),
        entities_list=entities,
        triples_list=triples,
        quiz_score=st.session_state.get("quiz_score", 0),
        quiz_total=quiz_total,
        student_notes=student_notes
    )

    if pdf_bytes:
        os.makedirs("static", exist_ok=True)
        with open("static/lab_report.pdf", "wb") as f:
            f.write(pdf_bytes)
        with open("lab_report.pdf", "wb") as f:
            f.write(pdf_bytes)

    st.divider()
    st.subheader("Download Official Lab Report (.pdf)")

    col_btn1, col_btn2 = st.columns(2)
    with col_btn1:
        st.link_button(
            "Open / Download PDF Document",
            url="/app/static/lab_report.pdf",
            type="primary",
            use_container_width=True
        )

    with col_btn2:
        st.download_button(
            label="Download lab_report.pdf",
            data=pdf_bytes,
            file_name="lab_report.pdf",
            mime="application/pdf",
            key="stream_pdf_btn",
            use_container_width=True
        )


# =============================================================================
# 7. MAIN ENTRYPOINT & NAVIGATION
# =============================================================================

def init_session_state():
    """Initializes Streamlit session state variables."""
    if "results" not in st.session_state:
        st.session_state["results"] = None
    if "input_text" not in st.session_state:
        st.session_state["input_text"] = SAMPLE_TEXTS["Example 1: Tech Companies & Partnerships"]
    if "prev_input_mode" not in st.session_state:
        st.session_state["prev_input_mode"] = "Sample Example"
    if "prev_sample_key" not in st.session_state:
        st.session_state["prev_sample_key"] = "Example 1: Tech Companies & Partnerships"
    if "prev_uploaded_file" not in st.session_state:
        st.session_state["prev_uploaded_file"] = None
    if "current_quiz_questions" not in st.session_state or len(st.session_state.get("current_quiz_questions", [])) != 5:
        st.session_state["current_quiz_questions"] = random.sample(QUIZ_QUESTIONS, 5)
    if "quiz_answers" not in st.session_state:
        st.session_state["quiz_answers"] = {}
    if "quiz_submitted" not in st.session_state:
        st.session_state["quiz_submitted"] = False
    if "quiz_score" not in st.session_state:
        st.session_state["quiz_score"] = 0
    if "student_info" not in st.session_state:
        st.session_state["student_info"] = {
            "name": "Student Name",
            "id": "EXP-001",
            "date": str(datetime.now().date())
        }
    if "student_notes" not in st.session_state:
        st.session_state["student_notes"] = ""


def main():
    st.set_page_config(
        page_title="KGIRS Virtual Lab — Relationship Extraction",
        page_icon="🕸️",
        layout="wide"
    )

    init_session_state()

    # Load spaCy NLP Model
    nlp, load_err = load_nlp_model("en_core_web_sm")
    if load_err:
        st.error(f"⚠️ **Model Loading Notice:** {load_err}")
        st.info("Run `python -m spacy download en_core_web_sm` to install the required model.")
        st.stop()

    # Native Header
    st.title("Experiment: Relationship Extraction from Text")
    st.caption("Knowledge Graph & Information Retrieval Systems (KGIRS) Virtual Laboratory")

    # Navigation Sidebar
    section = st.sidebar.radio(
        "Lab Navigator",
        options=["Theory", "Simulation", "Quiz", "Report Generation"]
    )

    st.sidebar.divider()
    st.sidebar.subheader("Progress Tracker")
    sim_status = "Done" if st.session_state.get("results") is not None else "Pending"
    st.sidebar.write(f"- **Simulation:** {sim_status}")
    if st.session_state.get("results") is not None:
        triples_count = len(st.session_state["results"].get("triples", []))
        st.sidebar.write(f"  - Triples: `{triples_count}`")

    quiz_total = len(st.session_state.get("current_quiz_questions", [])) or 5
    quiz_status = "Done" if st.session_state.get("quiz_submitted", False) else "Pending"
    st.sidebar.write(f"- **Quiz:** {quiz_status}")
    if st.session_state.get("quiz_submitted", False):
        st.sidebar.write(f"  - Score: `{st.session_state.get('quiz_score', 0)} / {quiz_total}`")

    # Section Dispatcher
    if section == "Theory":
        render_theory_section()
    elif section == "Simulation":
        render_simulation_section(nlp)
    elif section == "Quiz":
        render_quiz_section()
    elif section == "Report Generation":
        render_report_section()


if __name__ == "__main__":
    main()
